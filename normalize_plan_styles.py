from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Cm


SOURCE = r"C:\Users\jyy\Downloads\hospitaFXjyy\plan-for-review-revised.docx"
TARGET = r"C:\Users\jyy\Downloads\hospitaFXjyy\plan-for-review-revised-styled.docx"


HEADING_RENAMES = {
    "四、商业模式规划": "四、商业模式与推广路径",
    "4.3 销售渠道布局": "4.3 推广路径",
    "五、营销策略部署": "五、应用推广与服务策略",
    "5.2 具体获客策略": "5.2 试点拓展策略",
    "六、财务预测分析": "六、财务测算与分析",
    "未来五年项目收入预测": "6.1 项目收入测算",
    "运营成本结构剖析": "6.2 成本结构分析",
    "未来三年盈利预测展望": "6.3 收益展望",
    "七、核心团队": "七、团队与实施保障",
    "八、全面风险评估与系统化应对预案": "八、风险分析与应对",
    "8.2 融资与成果转化路径": "8.2 成果转化路径",
    "九、附录": "九、项目计划与附录",
    "9.1 项目关键里程碑计划表": "9.1 关键里程碑计划表",
}


def set_run_font(run, east_asia_name, ascii_name=None, size=None, bold=None):
    run.font.name = ascii_name or east_asia_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def replace_heading_texts(doc):
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in HEADING_RENAMES:
            p.text = HEADING_RENAMES[text]


def style_paragraph(p):
    text = p.text.strip()
    if not text:
        return

    if text == "蓉城医枢智慧门诊协同系统商业计划书":
        p.style = doc.styles["Title"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(12)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        for r in p.runs:
            set_run_font(r, "黑体", "Times New Roman", 20, True)
        return

    if p.style.name == "Heading 1":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.space_before = Pt(10)
        pf.space_after = Pt(6)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        for r in p.runs:
            set_run_font(r, "黑体", "Times New Roman", 16, True)
        return

    if p.style.name == "Heading 2":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.space_before = Pt(6)
        pf.space_after = Pt(3)
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        for r in p.runs:
            set_run_font(r, "黑体", "Times New Roman", 13, True)
        return

    if text == "结语":
        p.style = doc.styles["Heading 1"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.space_before = Pt(10)
        pf.space_after = Pt(6)
        for r in p.runs:
            set_run_font(r, "黑体", "Times New Roman", 16, True)
        return

    pf = p.paragraph_format
    pf.first_line_indent = Cm(0.74)
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        set_run_font(r, "宋体", "Times New Roman", 12, False)


def style_tables(doc):
    for table in doc.tables:
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    pf = p.paragraph_format
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    pf.first_line_indent = Cm(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        set_run_font(r, "宋体", "Times New Roman", 11 if row_idx == 0 else 10.5, row_idx == 0)


doc = Document(SOURCE)
replace_heading_texts(doc)

section = doc.sections[0]
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(2.8)
section.right_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(12)

for p in doc.paragraphs:
    style_paragraph(p)

style_tables(doc)
doc.save(TARGET)
print(TARGET)
